"""本体の紺を「主役の1か所」に絞る（rules/slide-rules.md G-6）。

見本デッキの書き出し後に build_pattern_deck.js から呼ばれる。任意のデッキにも使える：
    python3 scripts/one_navy.py deck.pptx [--keep P10=2,P51=1]

本体（メッセージ欄より下、注記・出所より上）で紺を使う要素を、近いもの（0.25インチ以内）どうしで「1か所」にまとめ、
主役（紺の塗りの面積が最も大きい1か所。同じなら上・左にあるもの）だけを残す。ほかの場所の紺は、
文字は文字色（#3F3F3F、太字はそのまま）、塗りは Gray2（#7F7F7F）、線は文字色に置き換える。
--keep で、型ごとに主役にする場所の番号（上から数えて0始まり）を指定できる。
"""
import sys, re
from pptx import Presentation
from lxml import etree

IN = 914400
NAVY, TEXT, GRAY2 = "000F78", "3F3F3F", "7F7F7F"
BODY_TOP, BODY_BOTTOM = 1.55, 6.8
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def _xml(sh):
    return etree.tostring(sh._element).decode()


def _uses_navy(sh):
    if f'val="{NAVY}"' in _xml(sh):
        return True
    return bool(getattr(sh, "has_chart", False) and sh.has_chart and f'val="{NAVY}"'.encode() in sh.chart.part.blob)


def _navy_area(sh):
    """紺の塗りの面積（インチ²）。グラフは紺の系列があれば図全体の1割と見なす。"""
    w, h = (sh.width or 0) / IN, (sh.height or 0) / IN
    if getattr(sh, "has_chart", False) and sh.has_chart:
        return w * h * 0.1
    sp = sh._element.find(f"{{http://schemas.openxmlformats.org/presentationml/2006/main}}spPr")
    if sp is not None:
        f = sp.find(f"{{{A}}}solidFill/{{{A}}}srgbClr")
        if f is not None and f.get("val", "").upper() == NAVY:
            return w * h
    return 0.01


def clusters(slide):
    items = []
    for sh in slide.shapes:
        y = (sh.top or 0) / IN
        if BODY_TOP <= y < BODY_BOTTOM and _uses_navy(sh):
            x, w, h = (sh.left or 0) / IN, (sh.width or 0) / IN, (sh.height or 0) / IN
            items.append((x, y, x + w, y + h, sh))
    groups = []
    for b in items:
        hit = [g for g in groups if any(b[0] - 0.25 <= q[2] and q[0] - 0.25 <= b[2] and b[1] - 0.25 <= q[3] and q[1] - 0.25 <= b[3] for q in g)]
        groups = [g for g in groups if g not in hit] + [[b] + [q for g in hit for q in g]]
    groups.sort(key=lambda g: (min(q[1] for q in g), min(q[0] for q in g)))
    return groups


def _demote(sh):
    el = sh._element
    for clr in el.iter(f"{{{A}}}srgbClr"):
        if clr.get("val", "").upper() != NAVY:
            continue
        parent = clr.getparent()          # solidFill
        owner = parent.getparent()        # spPr / ln / rPr / defRPr / tcPr ...
        tag = etree.QName(owner).localname
        clr.set("val", TEXT if tag in ("rPr", "defRPr", "endParaRPr", "ln", "lnL", "lnR", "lnT", "lnB") else GRAY2)
    if getattr(sh, "has_chart", False) and sh.has_chart:
        part = sh.chart.part
        part._element = etree.fromstring(part.blob.replace(f'val="{NAVY}"'.encode(), f'val="{GRAY2}"'.encode()))


def apply(path, keep=None, scope=re.compile(r"^(パターンID：P\d+|完成例)")):
    keep = keep or {}
    prs = Presentation(path)
    changed = 0
    for s in prs.slides:
        note = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        if not scope.match(note):
            continue
        pid = re.match(r"パターンID：(P\d+)", note)
        groups = clusters(s)
        if len(groups) < 2:
            continue
        k = keep.get(pid.group(1)) if pid else None
        main = groups[k] if k is not None and k < len(groups) else max(groups, key=lambda g: sum(_navy_area(q[4]) for q in g))
        for g in groups:
            if g is main:
                continue
            for q in g:
                _demote(q[4])
            changed += 1
    prs.save(path)
    return changed


if __name__ == "__main__":
    keep = {}
    if "--keep" in sys.argv:
        for kv in sys.argv[sys.argv.index("--keep") + 1].split(","):
            k, v = kv.split("=")
            keep[k] = int(v)
    print(f"one_navy: 紺を外した場所 {apply(sys.argv[1], keep)} か所")
