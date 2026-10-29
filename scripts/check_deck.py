import sys, re
from pptx import Presentation
from lxml import etree

# 指定色（rules/slide-rules.md B-3）：Navy・Blue1〜5・Gray1〜2・White・文字色・反対色
PALETTE = {"000F78", "3C64AA", "83A0D3", "72DFF0", "E0EEF8", "C1DDF1", "CCCCCC", "7F7F7F", "FFFFFF", "3F3F3F", "D73232"}
ACCENT_RED = "D73232"
PLACEHOLDER = re.compile(r"◯◯|○○|〇〇|XX|ＸＸ|ダミー|見本|lorem|TBD|テキストを入力|【(?!要確認)[^】]{1,15}】")
SMALL_OK = ("注", "出所", "（", "(", "縦軸", "横軸", "【補足", "補足", "※")

DARK = {"000F78", "7F7F7F", "3C64AA", "3F3F3F"}
CNS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart", "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

def zen_len(t):
    """全角換算の字数（半角英数・記号は0.55字）。2行に収まるかの目安に使う。"""
    return sum(0.55 if ord(c) < 0x2000 else 1 for c in t)

# ---- 版面の検査（rules/slide-rules.md H章。Jinba consulting-pptx-skill の check_layout を PPTX 用に簡略化したもの） ----
IN = 914400
NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
NS_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
BODY_TOP, BODY_BOTTOM = 1.9, 6.8          # 本体の範囲（インチ。スライドタイトルの線の下〜注記・出所の上）
TYPED_BULLET = re.compile(r"^[・•●■◆▪\-－‐]\s?")

def _rgb(el, path):
    c = el.find(path)
    return c.get("val").upper() if c is not None else None

def layout_checks(slide, i):
    """版面の検査結果を (区分, メッセージ) のリストで返す。"""
    out, bottom, boxes, filled_line, tint = [], 0.0, 0, 0, 0
    for sh in slide.shapes:
        top, h, w = (sh.top or 0) / IN, (sh.height or 0) / IN, (sh.width or 0) / IN
        if getattr(sh, "has_table", False) and sh.has_table:
            h = sum(r.height for r in sh.table.rows) / IN
            # H-5 最終行の下に罫線を引かない
            last = sh.table._tbl.tr_lst[-1]
            if any(tc.find(f"{NS_A}tcPr/{NS_A}lnB/{NS_A}solidFill") is not None for tc in last.tc_lst):
                out.append(("WARN", "表の最終行の下に罫線がある（H-5）"))
        empty_text = sh.has_text_frame and not sh.text_frame.text.strip()
        if BODY_TOP - 0.05 <= top < BODY_BOTTOM and not (empty_text and h < 0.05):
            bottom = max(bottom, min(top + h, BODY_BOTTOM))
        pr = sh._element.find(f"{NS_P}spPr")
        if pr is not None:
            geom = pr.find(f"{NS_A}prstGeom")
            fill = _rgb(pr, f"{NS_A}solidFill/{NS_A}srgbClr")
            ln = pr.find(f"{NS_A}ln")
            lcol = _rgb(ln, f"{NS_A}solidFill/{NS_A}srgbClr") if ln is not None else None
            if geom is not None and geom.get("prst") in ("rect", "chevron", "homePlate"):
                if fill and fill != "FFFFFF" and lcol and lcol != "FFFFFF": filled_line += 1
                if geom.get("prst") == "rect" and lcol and lcol != "FFFFFF" and fill in (None, "FFFFFF") and w > 0.9 and h > 0.35: boxes += 1
                if fill == "C1DDF1" and w > 0.6 and h > 0.3: tint += 1
        # H-6 記号の直打ち、H-8 文字のはみ出し（見積もり）
        if sh.has_text_frame and sh.text_frame.text.strip() and w > 0.3 and h > 0.15 and top > 0.2:
            need = 0.0
            for para in sh.text_frame.paragraphs:
                t = "".join(r.text for r in para.runs)
                if TYPED_BULLET.match(t) and len(t) > 3:
                    out.append(("WARN", f"箇条書きの記号を文字で打っている（H-6）：{t[:16]}"))
                size = max([r.font.size.pt for r in para.runs if r.font.size] or [14])
                per_line = max(1.0, (w * 72) / size)
                need += max(1, -(-zen_len(t) // per_line)) * size * 1.2 / 72
            if need > h + 0.15:
                out.append(("WARN", f"文字が枠からはみ出す恐れ（必要{need:.2f}in／枠{h:.2f}in、H-8）：{sh.text_frame.text.strip()[:16]}"))
    if filled_line: out.append(("FAIL", f"塗りのある図形に枠線が重なっている（{filled_line}個、H-3）"))
    if boxes >= 4: out.append(("WARN", f"線で囲んだ箱が{boxes}個ある。器は見出し＋線で作り、囲むのは立てたい1〜2個だけ（H-1。ツリー・フロー・体制図は可）"))
    if tint >= 2: out.append(("WARN", f"淡い紺（Blue5）の面が{tint}か所ある。注目の面は1スライド1か所（H-4）"))
    if bottom:
        empty = 1 - (bottom - BODY_TOP) / (BODY_BOTTOM - BODY_TOP)
        if empty > 0.4: out.append(("FAIL", f"本体の下{empty:.0%}が空いている。中身を足すか型を変える（引き伸ばしは不可、H-7）"))
        elif empty > 0.3: out.append(("WARN", f"本体の下{empty:.0%}が空いている（H-7）"))
    return out

# ---- 洗練の検査（rules/slide-rules.md G-6〜G-12「洗練の核心」） ----
NAVY = "000F78"
DECIMAL_BIG = re.compile(r"(?<![\d.,])\d{1,3}(?:,\d{3})+\.\d|(?<![\d.,])\d{3,}\.\d")

def _uses_navy(sh):
    x = etree.tostring(sh._element).decode()
    if f'val="{NAVY}"' in x: return True
    if getattr(sh, "has_chart", False) and sh.has_chart:
        return f'val="{NAVY}"' in sh.chart.part.blob.decode("utf-8", "ignore")
    return False

def _clusters(boxes, gap=0.25):
    """近い（gap以内）箱をまとめて1か所と数える。"""
    groups = []
    for b in boxes:
        hit = [g for g in groups if any(b[0] - gap <= c[2] and c[0] - gap <= b[2] and b[1] - gap <= c[3] and c[1] - gap <= b[3] for c in g)]
        merged = [b] + [c for g in hit for c in g]
        groups = [g for g in groups if g not in hit] + [merged]
    return len(groups)

def polish_checks(slide, navy_check=True):
    out, navy, note_lines, note_chars = [], [], 0, 0
    for sh in slide.shapes:
        x, y, w, h = (sh.left or 0) / IN, (sh.top or 0) / IN, (sh.width or 0) / IN, (sh.height or 0) / IN
        text = sh.text_frame.text.strip() if sh.has_text_frame else ""
        # G-6 本体の紺（メッセージ欄・ページ番号を除く。近い要素は1か所と数える）
        if BODY_TOP - 0.35 <= y < BODY_BOTTOM and _uses_navy(sh):
            navy.append((x, y, x + w, y + h))
        # G-10 注記・出所の量
        if y >= BODY_BOTTOM - 0.05 and text.startswith(("注", "出所", "資料")):
            note_lines += len([l for l in text.splitlines() if l.strip()]); note_chars += len(text)
        # G-8 100以上の値の小数
        if BODY_TOP - 0.35 <= y < BODY_BOTTOM and DECIMAL_BIG.search(text):
            out.append(("WARN", f"100以上の値に小数が付いている（G-8）：{DECIMAL_BIG.search(text).group(0)}"))
        if getattr(sh, "has_chart", False) and sh.has_chart:
            root = etree.fromstring(sh.chart.part.blob)
            fmts = [f.get("formatCode", "") for f in root.iter("{%s}numFmt" % CNS["c"])]
            vals = [float(v.text) for v in root.findall(".//c:val//c:v", CNS) if re.fullmatch(r"-?\d+(\.\d+)?", v.text or "")]
            if any(re.search(r"0\.0", f) for f in fmts) and vals and max(abs(v) for v in vals) >= 100:
                out.append(("WARN", "100以上の値を含むグラフのラベルが小数の書式（G-8）"))
        # G-12 表のセルの折り返し（3行以上）
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for ci, cell in enumerate(row.cells):
                    cw = sh.table.columns[ci].width / IN
                    lines = 0
                    for para in cell.text_frame.paragraphs:
                        t = "".join(r.text for r in para.runs)
                        size = max([r.font.size.pt for r in para.runs if r.font.size] or [14])
                        lines += max(1, -(-zen_len(t) // max(1.0, ((cw - 0.2) * 72) / size))) if t else 0
                    if lines >= 3:
                        out.append(("WARN", f"表のセルが{int(lines)}行に折り返す（G-12）：{cell.text_frame.text.strip()[:12]}")); break
    n = _clusters(navy)
    if navy_check and n >= 2: out.append(("WARN", f"本体の紺が{n}か所ある。紺はメッセージが指す1か所に絞る（G-6）"))
    if note_lines >= 3: out.append(("WARN", f"注記・出所が{note_lines}行ある。2行までに絞り、定義は本体の見出しへ（G-10）"))
    elif note_chars > 150: out.append(("WARN", f"注記・出所が{note_chars}字ある。出所は略記し、定義の説明を入れない（G-10）"))
    return out

# ---- 中身の厚みの検査（rules/slide-rules.md I-1：同じ数字を描き方を変えて繰り返さない） ----
KEYNUM = re.compile(r"(?<![\d.,])\d{1,3}(?:,\d{3})+(?:\.\d+)?(?![\d])|(?<![\d.,])\d+\.\d+(?![\d])|(?<![\d.,])\d{3,}(?![\d.,])")

def repeated_numbers(prs, limit=4):
    """本体（メッセージ・注記を除く）に出る3桁以上・小数の数字が、limit枚以上のページに出ていればWARN。年（1900〜2099）は除く。"""
    where = {}
    for i, s in enumerate(prs.slides, 1):
        if i == 1: continue
        found = set()
        for sh in s.shapes:
            y = (sh.top or 0) / IN
            if not (BODY_TOP - 0.35 <= y < BODY_BOTTOM) or not sh.has_text_frame: continue
            for n in KEYNUM.findall(sh.text_frame.text):
                v = n.replace(",", "")
                if re.fullmatch(r"(19|20)\d\d", v): continue
                found.add(v)
        for v in found: where.setdefault(v, []).append(i)
    out = []
    for v, ps in sorted(where.items(), key=lambda kv: -len(kv[1])):
        if len(ps) >= limit:
            out.append(f"同じ数字「{v}」が{len(ps)}ページに出ている（p{', p'.join(map(str, ps))}）。描き方を変えて繰り返さず、各ページで新しい事実を示す（I-1）")
    return out[:5]

# ---- ガイド線の検査（rules/slide-rules.md B-1「ガイド線」） ----
DEFAULT_GUIDES_CM = {"title_top": 8.80, "body_top": 5.40, "left": -13.20, "right": 13.20}

def read_guides(prs):
    """スライドマスターのガイド線（1/8pt単位）を読み、インチで返す。無ければルールの既定値から計算する。"""
    sw, sh = prs.slide_width / IN, prs.slide_height / IN
    g = {"title_top": sh / 2 - DEFAULT_GUIDES_CM["title_top"] / 2.54, "body_top": sh / 2 - DEFAULT_GUIDES_CM["body_top"] / 2.54,
         "left": sw / 2 + DEFAULT_GUIDES_CM["left"] / 2.54, "right": sw / 2 + DEFAULT_GUIDES_CM["right"] / 2.54, "found": False}
    xml = etree.tostring(prs.slide_masters[0]._element).decode()
    horz = sorted(int(m) / 576 for m in re.findall(r'<p15:guide[^>]*orient="horz"[^>]*pos="(\d+)"', xml))
    vert = sorted(int(m) / 576 for m in re.findall(r'<p15:guide(?![^>]*orient="horz")[^>]*pos="(\d+)"', xml))
    if len(horz) >= 2 and len(vert) >= 2:
        g.update(title_top=horz[0], body_top=horz[1], left=vert[0], right=vert[-1], found=True)
    return g

def guide_checks(slide, g):
    out, tol = [], 0.03
    tops = []
    for sh in slide.shapes:
        x, y, w = (sh.left or 0) / IN, (sh.top or 0) / IN, (sh.width or 0) / IN
        if w <= 0: continue
        if x < g["left"] - tol or x + w > g["right"] + tol:
            label = sh.text_frame.text.strip()[:12] if sh.has_text_frame else sh.shape_type
            out.append(("FAIL", f"ガイド線③④（左端・右端）の外にはみ出している：{label}"))
        tops.append(y)
    if not any(abs(t - g["title_top"]) <= 0.04 for t in tops):
        out.append(("WARN", "メッセージブロックの上端がガイド線①にそろっていない"))
    if not any(abs(t - g["body_top"]) <= 0.07 for t in tops):
        out.append(("WARN", "ボディ（本体の見出し）の上端がガイド線②にそろっていない"))
    return out

def label_contrast(chart_part):
    """グラフのデータラベルの色と系列の塗りの対比を点検する（濃い塗りの内側ラベルは白、淡い塗りは黒）。"""
    root = etree.fromstring(chart_part.blob)
    bad = []
    for ser in root.iter("{%s}ser" % CNS["c"]):
        fill = ser.find("c:spPr/a:solidFill/a:srgbClr", CNS)
        d = ser.find("c:dLbls", CNS)
        if fill is None or d is None or d.find("c:showVal", CNS) is None or d.find("c:showVal", CNS).get("val") != "1":
            continue
        pos = d.find("c:dLblPos", CNS)
        if pos is None:
            pos = ser.getparent().find("c:dLbls/c:dLblPos", CNS)
        if pos is None:
            g = ser.getparent().find("c:grouping", CNS)
            inside = g is not None and g.get("val") in ("stacked", "percentStacked")
        else:
            inside = pos.get("val") in ("ctr", "inEnd", "inBase")
        clr = d.find(".//c:txPr//a:defRPr/a:solidFill/a:srgbClr", CNS)
        lc = (clr.get("val") if clr is not None else "000000").upper()
        fc = fill.get("val").upper()
        if inside and fc in DARK and lc != "FFFFFF": bad.append(f"濃い塗り#{fc}に#{lc}のラベル")
        if (not inside or fc not in DARK) and lc == "FFFFFF": bad.append(f"淡い塗り・外側に白ラベル（#{fc}）")
    return bad

def main(path, sample=False):
    prs = Presentation(path)
    guides = read_guides(prs)
    H = prs.slide_height
    fails, warns = [], []
    if not guides["found"]:
        warns.append((0, "スライドマスターにガイド線がない（python3 scripts/add_guides.py deck.pptx で入れる。検査は既定値で行う）"))
    for i, s in enumerate(prs.slides, 1):
        xml = etree.tostring(s._element).decode()
        # 色：パレット外
        for c in set(re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', xml)):
            c = c.upper()
            if c == ACCENT_RED: warns.append((i, "反対色(#D73232)を使用：悪化・未達そのものがメッセージか確認"))
            if c not in PALETTE: fails.append((i, f"パレット外の色 #{c}"))
        # 装飾
        if 'prst="roundRect"' in xml: fails.append((i, "角丸四角形"))
        if "outerShdw" in xml or "gradFill" in xml or "sp3d" in xml: fails.append((i, "影・グラデーション・3D"))
        msg, body_chars = None, 0
        for sh in s.shapes:
            if getattr(sh, "has_chart", False) and sh.has_chart:
                for b in label_contrast(sh.chart.part): fails.append((i, f"データラベルの対比：{b}"))
            if not sh.has_text_frame: continue
            t = sh.text_frame.text.strip()
            if not t: continue
            top = (sh.top or 0) / H
            if not sample and PLACEHOLDER.search(t): fails.append((i, f"仮の文字が残存：{t[:20]}"))
            if "【要確認】" in t: warns.append((i, "【要確認】が残存"))
            sizes = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
            if sizes and max(sizes) >= 19.5 and top < 0.2:
                msg = t.replace("\n", "")
                if any(r.font.bold for p in sh.text_frame.paragraphs for r in p.runs if r.text.strip()):
                    warns.append((i, "メッセージが太字。標準の太さにする（B-2）"))
                continue
            zone_ok = top > 0.88 or top < 0.08 or t.startswith(SMALL_OK) or len(t) <= 8
            # B-2：本文・読み取り14pt、図の文字・補足12pt、注記9pt。12pt未満は注記・出所・章名だけ
            bad = [r for p in sh.text_frame.paragraphs for r in p.runs
                   if r.font.size and r.font.size.pt < 12 and r.text.strip() and not r.text.strip().startswith(SMALL_OK)]
            if bad and not zone_ok:
                fails.append((i, f"12pt未満（{bad[0].font.size.pt:g}pt）：{bad[0].text.strip()[:20]}"))
            if 0.25 < top < 0.9 and not t.startswith(("注", "出所")): body_chars += len(t)
        if i > 1 and msg:
            n = len(msg)
            if n < 55: warns.append((i, f"メッセージが短い（{n}字）：2行を埋めているか"))
            if n > 80: fails.append((i, f"メッセージが長い（{n}字）：3行になる恐れ"))
            elif zen_len(msg) > 69: warns.append((i, f"メッセージが全角換算{zen_len(msg):.0f}字：3行になる恐れ（66字以内が目安）"))
            if re.match(r"^(この|その|これ|それ)", msg): fails.append((i, "メッセージが他ページを指示語で参照"))
            if re.match(r"^[^。]{1,12}[：:]", msg): fails.append((i, "メッセージに「見出し：」の前置き"))
            if re.search(r"(です|ます)。?$", msg): fails.append((i, "メッセージが「です・ます」調"))
            if re.search(r"\d[\d,.]*\s?[万億千]?(円|人|件|t|トン|回|本)\s?[／/]\s?(年|人|日|月)", msg): warns.append((i, "メッセージの単位を記号で書いている。「年142億円」のように和文で書く（G-8）"))
        if body_chars > 450: warns.append((i, f"本体の文字が多い（約{body_chars}字）：450字を超えたら分ける（C-1）"))
        if i > 1 and msg:
            for tag, m in layout_checks(s, i):
                (fails if tag == "FAIL" else warns).append((i, m))
            for tag, m in guide_checks(s, guides):
                (fails if tag == "FAIL" else warns).append((i, m))
            notes0 = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
            # 見本デッキの記号集・アイコン集・パーツ集・テンプレート集は部品の一覧なので、紺の数は数えない
            for tag, m in polish_checks(s, navy_check=not (sample and not re.match(r"(パターンID：P\d|完成例)", notes0))):
                (fails if tag == "FAIL" else warns).append((i, m))
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        if i > 1 and not re.search(r"E-?\d|エビデンス|P\d\d", notes): warns.append((i, "ノートにエビデンスID・パターンIDなし"))
    if not sample:
        for m in repeated_numbers(prs):
            warns.append((0, m))
    for tag, items in (("FAIL", fails), ("WARN", warns)):
        for p, m in items: print(f"{tag}  p{p:02d}  {m}")
    print(f"--- FAIL {len(fails)} / WARN {len(warns)}")
    return 1 if fails else 0

if __name__ == "__main__":
    sys.exit(main(sys.argv[1], "--sample" in sys.argv))