import sys, re
from pptx import Presentation
from lxml import etree

PALETTE = {"000F78", "000000", "595959", "7F7F7F", "BFBFBF", "E7E6E6", "F2F2F2", "FFFFFF", "D9D9D9", "404040", "A6A6A6"}
PLACEHOLDER = re.compile(r"◯◯|○○|〇〇|XX|ＸＸ|ダミー|見本|lorem|TBD|テキストを入力|【(?!要確認)[^】]{1,15}】")
SMALL_OK = ("注", "出所", "（", "(", "縦軸", "横軸", "【補足", "補足", "※")

def main(path, sample=False):
    prs = Presentation(path)
    H = prs.slide_height
    fails, warns = [], []
    for i, s in enumerate(prs.slides, 1):
        xml = etree.tostring(s._element).decode()
        # 色：パレット外
        for c in set(re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', xml)):
            c = c.upper()
            if c == "C00000": warns.append((i, "赤(#C00000)を使用：悪化そのものがメッセージか確認"))
            elif c not in PALETTE: fails.append((i, f"パレット外の色 #{c}"))
        # 装飾
        if 'prst="roundRect"' in xml: fails.append((i, "角丸四角形"))
        if "outerShdw" in xml or "gradFill" in xml or "sp3d" in xml: fails.append((i, "影・グラデーション・3D"))
        msg, body_chars = None, 0
        for sh in s.shapes:
            if not sh.has_text_frame: continue
            t = sh.text_frame.text.strip()
            if not t: continue
            top = (sh.top or 0) / H
            if not sample and PLACEHOLDER.search(t): fails.append((i, f"仮の文字が残存：{t[:20]}"))
            if "【要確認】" in t: warns.append((i, "【要確認】が残存"))
            sizes = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
            if sizes and max(sizes) >= 19.5 and top < 0.2: msg = t.replace("\n", ""); continue
            zone_ok = top > 0.88 or top < 0.08 or t.startswith(SMALL_OK) or len(t) <= 8
            bad = [r for p in sh.text_frame.paragraphs for r in p.runs
                   if r.font.size and r.font.size.pt < 14 and r.text.strip() and not r.text.strip().startswith(SMALL_OK)]
            if bad and not zone_ok:
                fails.append((i, f"14pt未満（{bad[0].font.size.pt:g}pt）：{bad[0].text.strip()[:20]}"))
            if 0.25 < top < 0.9 and not t.startswith(("注", "出所")): body_chars += len(t)
        if i > 1 and msg:
            n = len(msg)
            if n < 55: warns.append((i, f"メッセージが短い（{n}字）：2行を埋めているか"))
            if n > 80: fails.append((i, f"メッセージが長い（{n}字）：3行になる恐れ"))
            if re.match(r"^(この|その|これ|それ)", msg): fails.append((i, "メッセージが他ページを指示語で参照"))
            if re.match(r"^[^。]{1,12}[：:]", msg): fails.append((i, "メッセージに「見出し：」の前置き"))
            if re.search(r"(です|ます)。?$", msg): fails.append((i, "メッセージが「です・ます」調"))
        if body_chars > 200: warns.append((i, f"本体の文字が多い（約{body_chars}字）"))
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        if i > 1 and not re.search(r"E-?\d|エビデンス|P\d\d", notes): warns.append((i, "ノートにエビデンスID・パターンIDなし"))
    for tag, items in (("FAIL", fails), ("WARN", warns)):
        for p, m in items: print(f"{tag}  p{p:02d}  {m}")
    print(f"--- FAIL {len(fails)} / WARN {len(warns)}")
    return 1 if fails else 0

if __name__ == "__main__":
    sys.exit(main(sys.argv[1], "--sample" in sys.argv))